"""语料读取：把 txt/md/pdf/html 统一转成纯文本。

返回：{ 文档名: 纯文本 }。其余格式留扩展位。
"""
import os
import re


def read_txt(path):
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def read_md(path):
    return read_txt(path)


def read_html(path):
    from html.parser import HTMLParser

    class _P(HTMLParser):
        def __init__(self):
            super().__init__()
            self.text = []
        def handle_data(self, d):
            if d.strip():
                self.text.append(d)

    p = _P()
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        p.feed(f.read())
    return "\n".join(p.text)


def read_pdf(path):
    """可选依赖 pypdf；未安装时给出提示。"""
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError("解析 PDF 需要 pypdf，请 `pip install pypdf`")
    reader = PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def read_docx(path):
    """读 .docx 的段落 + 表格文本（培养方案/规定常以表格呈现，很重要）。"""
    import docx
    d = docx.Document(path)

    def _iter_text(paragraphs):
        for p in paragraphs:
            t = p.text.strip()
            if t:
                yield t

    parts = list(_iter_text(d.paragraphs))
    # 表格：逐行把单元格文本拼起来，便于后续分块
    for tbl in d.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


_READERS = {
    ".txt": read_txt,
    ".md": read_md,
    ".markdown": read_md,
    ".html": read_html,
    ".htm": read_html,
    ".pdf": read_pdf,
    ".docx": read_docx,
}


def read_all_docs(corpus_dir, skip_thumb=True):
    """扫描 corpus_dir 下所有支持格式，返回 {文档名: 文本}。"""
    docs = {}
    if not os.path.isdir(corpus_dir):
        raise FileNotFoundError(f"语料目录不存在: {corpus_dir}")
    for root, _, files in os.walk(corpus_dir):
        for fn in files:
            ext = os.path.splitext(fn)[1].lower()
            if ext not in _READERS:
                continue
            if skip_thumb and fn.startswith("~$"):
                continue
            path = os.path.join(root, fn)
            rel = os.path.relpath(path, corpus_dir)
            try:
                text = _READERS[ext](path)
                # 清理多余空行，便于分块
                cleaned = re.sub(r"\n{3,}", "\n\n", text).strip()
                if cleaned:
                    docs[rel] = cleaned
            except Exception as e:
                print(f"[loader] 跳过 {rel}: {e}")
    return docs